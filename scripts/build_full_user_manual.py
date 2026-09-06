import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import sys, os

# Import helpers
sys.path.append(os.path.dirname(__file__))
from docx_helpers import set_cell_background, set_cell_margins, set_cell_border, add_callout, style_table

def build_manual():
    doc = docx.Document()
    
    # Page setup (Letter, margins 1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        
        # Header / Footer setup
        footer = section.footer
        p_f = footer.paragraphs[0]
        p_f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = p_f.add_run("MERCAPLENO | Manual de Usuario v2.0.0")
        f_run.font.name = "Segoe UI"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(148, 163, 184)
        
        header = section.header
        p_h = header.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h_run = p_h.add_run("Sistema de Gestión de Ventas e Inventario")
        h_run.font.name = "Segoe UI"
        h_run.font.size = Pt(8.5)
        h_run.font.color.rgb = RGBColor(148, 163, 184)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # PORTADA
    # -------------------------------------------------------------
    p_pre = doc.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(36)
    p_pre.paragraph_format.space_after = Pt(12)
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_tag = p_pre.add_run("PROYECTO FORMATIVO / SISTEMA INTEGRADO DE GESTIÓN")
    r_tag.font.name = "Segoe UI"
    r_tag.font.size = Pt(10)
    r_tag.font.bold = True
    r_tag.font.color.rgb = RGBColor(37, 99, 235)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("MANUAL DE USUARIO")
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 58, 138)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(4)
    p_sub.paragraph_format.space_after = Pt(28)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Sistema Web y Móvil de Gestión de Ventas, Inventario y Catálogo Comercial")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Decorative line
    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_after = Pt(36)
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_line = p_line.add_run("—" * 32)
    r_line.font.color.rgb = RGBColor(203, 213, 225)

    # Metadata table on cover
    cover_table = doc.add_table(rows=5, cols=2)
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Nombre del Software:", "MERCAPLENO"),
        ("Versión del Sistema:", "2.0.0 (Migración PostgreSQL 16 & Docker)"),
        ("Fecha de Publicación:", "Septiembre de 2026"),
        ("Autores y Desarrollo:", "Jeferson Jair Bernal Castro\nSebastian Rivera\nDylan Sneider Rivera Mora\nJuan Perez"),
        ("Entorno de Ejecución:", "Arquitectura Web React 18 / NestJS / PostgreSQL 16")
    ]
    for i, (k, v) in enumerate(meta_data):
        c0 = cover_table.cell(i, 0)
        c1 = cover_table.cell(i, 1)
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        set_cell_border(c0, bottom={"val": "single", "sz": "4", "color": "E2E8F0"},
                            top={"val": "single", "sz": "4", "color": "E2E8F0"},
                            left={"val": "single", "sz": "4", "color": "E2E8F0"},
                            right={"val": "single", "sz": "4", "color": "E2E8F0"})
        set_cell_border(c1, bottom={"val": "single", "sz": "4", "color": "E2E8F0"},
                            top={"val": "single", "sz": "4", "color": "E2E8F0"},
                            left={"val": "single", "sz": "4", "color": "E2E8F0"},
                            right={"val": "single", "sz": "4", "color": "E2E8F0"})
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = RGBColor(30, 41, 59)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_page_break()

    # -------------------------------------------------------------
    # CONTROL DE VERSIONES Y APROBACIONES
    # -------------------------------------------------------------
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(12)
    p_h1.paragraph_format.space_after = Pt(8)
    r = p_h1.add_run("Control de Versiones del Documento")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(30, 58, 138)

    t_vers = doc.add_table(rows=3, cols=4)
    t_vers_data = [
        ["Versión", "Fecha", "Autores", "Descripción del Cambio"],
        ["1.0.0", "02/09/2026", "Jair Bernal, Sebastian Rivera", "Borrador inicial del manual de usuario funcional."],
        ["2.0.0", "03/09/2026", "Jeferson Castro, Equipo Mercapleno", "Versión profesional completa: integración de módulo de carrito (cart), arquitectura Docker, PostgreSQL 16, catálogo depurado, procedimientos paso a paso y solución de problemas."]
    ]
    for r_idx, row in enumerate(t_vers_data):
        for c_idx, val in enumerate(row):
            t_vers.cell(r_idx, c_idx).paragraphs[0].add_run(val)
    style_table(t_vers)

    p_aprob = doc.add_paragraph()
    p_aprob.paragraph_format.space_before = Pt(18)
    p_aprob.paragraph_format.space_after = Pt(8)
    r = p_aprob.add_run("Firmas y Aprobaciones del Sistema")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor(30, 58, 138)

    t_apr = doc.add_table(rows=5, cols=4)
    t_apr_data = [
        ["Rol / Cargo", "Nombre Completo", "Firma / Aprobación", "Fecha"],
        ["Product Owner", "Jeferson Jair Bernal Castro", "Aprobado (Digital)", "03/09/2026"],
        ["Scrum Master", "Sebastian Rivera", "Aprobado (Digital)", "03/09/2026"],
        ["Líder de Desarrollo", "Dylan Sneider Rivera Mora", "Aprobado (Digital)", "03/09/2026"],
        ["Control de Calidad (QA)", "Juan Perez", "Aprobado (Digital)", "03/09/2026"]
    ]
    for r_idx, row in enumerate(t_apr_data):
        for c_idx, val in enumerate(row):
            t_apr.cell(r_idx, c_idx).paragraphs[0].add_run(val)
    style_table(t_apr)

    # -------------------------------------------------------------
    # TABLA DE CONTENIDO (RESUMEN ESTRUCTURAL)
    # -------------------------------------------------------------
    p_toc_h = doc.add_paragraph()
    p_toc_h.paragraph_format.space_before = Pt(24)
    p_toc_h.paragraph_format.space_after = Pt(8)
    r = p_toc_h.add_run("Tabla de Contenido General")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(30, 58, 138)

    toc_items = [
        ("1. Introducción y Marco General", "Propósito, objetivos, alcance y audiencia."),
        ("2. Requisitos del Sistema y Arquitectura", "Requisitos técnicos, URLs de acceso y matriz de permisos por rol."),
        ("3. Módulo 1: Seguridad, Registro y Autenticación (EP-001)", "Registro de clientes, verificación por correo (OTP), login 2FA y recuperación de clave."),
        ("4. Módulo 2: Catálogo de Productos y Tienda Virtual (EP-003)", "Búsqueda en tiempo real, filtros por categoría, consulta de existencias."),
        ("5. Módulo 3: Carrito de Compras y Proceso de Venta (EP-003)", "Gestión de ítems, control de stock, cálculo de IVA (19%), métodos de pago y ticket digital."),
        ("6. Módulo 4: Dashboard y Panel de Operaciones", "Panel de administración y accesos rápidos para roles 1 y 2."),
        ("7. Módulo 5: Gestión de Inventario y Registro de Movimientos (EP-004)", "Entradas por compras, salidas, devoluciones y semáforo de stock bajo."),
        ("8. Módulo 6: Administración del Catálogo de Productos (EP-004)", "Alta, edición, fijación de precios, estados y carga de imágenes."),
        ("9. Módulo 7: Gestión de Proveedores", "Directorio, habilitación, deshabilitación y actualización de contactos."),
        ("10. Módulo 8: Administración de Usuarios y Roles (EP-002)", "Creación de cuentas de empleado/admin, asignación de roles y control de acceso."),
        ("11. Módulo 9: Reportes, Analítica y Exportación PDF (EP-005)", "KPIs mensuales, Top productos más vendidos y descarga de reportes ejecutivos en PDF."),
        ("12. Guía de Solución de Problemas (Troubleshooting) y Preguntas Frecuentes", "Matriz de incidentes comunes y respuestas a preguntas frecuentes."),
        ("13. Glosario Técnico y de Términos del Negocio", "Definición de conceptos clave del aplicativo.")
    ]

    t_toc = doc.add_table(rows=len(toc_items)+1, cols=2)
    t_toc.cell(0, 0).paragraphs[0].add_run("Capítulo / Módulo")
    t_toc.cell(0, 1).paragraphs[0].add_run("Descripción del Contenido")
    for idx, (cap, desc) in enumerate(toc_items, start=1):
        t_toc.cell(idx, 0).paragraphs[0].add_run(cap)
        t_toc.cell(idx, 1).paragraphs[0].add_run(desc)
    style_table(t_toc)

    doc.add_page_break()

    # -------------------------------------------------------------
    # CAPÍTULO 1: INTRODUCCIÓN Y MARCO GENERAL
    # -------------------------------------------------------------
    def add_section_header(title, level=1):
        p = doc.add_paragraph()
        if level == 1:
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(title)
            r.bold = True
            r.font.size = Pt(16)
            r.font.color.rgb = RGBColor(30, 58, 138)
        elif level == 2:
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(title)
            r.bold = True
            r.font.size = Pt(13)
            r.font.color.rgb = RGBColor(37, 99, 235)
        else:
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(title)
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(51, 65, 85)

    add_section_header("1. Introducción y Marco General", level=1)
    
    doc.add_paragraph(
        "El presente Manual de Usuario constituye la guía oficial y estandarizada para la operación, "
        "administración y aprovechamiento óptimo de MERCAPLENO, una solución tecnológica integral de software "
        "diseñada para modernizar y optimizar los flujos de venta comercial, control riguroso de inventarios, "
        "gestión de compras con proveedores y generación de analítica de negocio en tiempo real."
    )
    
    add_section_header("1.1. Propósito del Documento", level=2)
    doc.add_paragraph(
        "Este documento tiene como finalidad instruir de manera secuencial y didáctica a los diferentes actores "
        "que interactúan con la plataforma, garantizando que cada procedimiento operativo se ejecute cumpliendo "
        "las reglas de negocio, los estándares de seguridad de la información y la integridad de los datos transaccionales."
    )

    add_section_header("1.2. Objetivos del Sistema", level=2)
    p_obj = doc.add_paragraph()
    p_obj.add_run("• Objetivo General: ").bold = True
    p_obj.add_run(
        "Proporcionar una herramienta robusta y escalable que centralice la administración de productos, "
        "la gestión de inventario multietapa, el proceso de autoservicio de ventas mediante carrito de compras y "
        "la toma de decisiones estratégicas mediante reportería consolidada."
    )
    
    p_objs = doc.add_paragraph()
    p_objs.add_run("• Objetivos Específicos:\n").bold = True
    p_objs.add_run("  1. Guiar a los usuarios en los procesos de registro, autenticación segura y doble factor (2FA).\n")
    p_objs.add_run("  2. Detallar el flujo de compra para clientes, desde la exploración del catálogo hasta la emisión del comprobante digital.\n")
    p_objs.add_run("  3. Estandarizar el registro de movimientos de inventario (entradas, salidas y devoluciones) para empleados y administradores.\n")
    p_objs.add_run("  4. Facilitar la administración de usuarios, asignación de roles y control de proveedores.\n")
    p_objs.add_run("  5. Instruir en la generación y exportación formal de reportes ejecutivos en PDF.")

    add_section_header("1.3. Alcance y Audiencia", level=2)
    doc.add_paragraph(
        "El alcance del manual cubre la totalidad de los módulos desarrollados en el frontend web (React + Vite) "
        "y los servicios backend (NestJS con PostgreSQL 16). La audiencia comprende tres perfiles principales: "
        "Administrador General, Empleado Operativo y Cliente Final."
    )

    add_callout(doc, 
                "Las instrucciones y capturas descritas en este manual corresponden a la versión 2.0.0 de Mercapleno. "
                "Cualquier modificación en los flujos transaccionales será reflejada en actualizaciones numeradas de este documento.",
                title="CONVENCIONES Y VIGENCIA", callout_type="NOTE")

    # -------------------------------------------------------------
    # CAPÍTULO 2: REQUISITOS Y ARQUITECTURA
    # -------------------------------------------------------------
    add_section_header("2. Requisitos del Sistema y Arquitectura de Operación", level=1)
    
    doc.add_paragraph(
        "Mercapleno está construido bajo una arquitectura modular desacoplada basada en microservicios y contenedores Docker. "
        "A continuación se detallan los requerimientos técnicos y las vías de acceso oficiales."
    )

    add_section_header("2.1. Requisitos Técnicos del Cliente", level=2)
    t_req = doc.add_table(rows=5, cols=2)
    t_req_data = [
        ["Componente", "Especificación Recomendada"],
        ["Dispositivo", "Computador de escritorio, portátil, tablet o smartphone con conexión a red."],
        ["Navegador Web", "Google Chrome (v110+), Mozilla Firefox (v110+), Microsoft Edge (v110+), Safari (v16+)."],
        ["Resolución de Pantalla", "Mínimo 1280 x 720 píxeles (Diseño totalmente responsivo adaptable a móviles)."],
        ["Conectividad", "Conexión a red local o Internet con ancho de banda mínimo de 2 Mbps."]
    ]
    for r_idx, row in enumerate(t_req_data):
        for c_idx, val in enumerate(row):
            t_req.cell(r_idx, c_idx).paragraphs[0].add_run(val)
    style_table(t_req)

    add_section_header("2.2. URLs de Acceso y Puertos de Despliegue", level=2)
    doc.add_paragraph(
        "El sistema opera bajo los siguientes puntos de enlace oficiales:"
    )
    p_urls = doc.add_paragraph()
    p_urls.add_run("• Aplicación Web (Frontend): ").bold = True
    p_urls.add_run("http://localhost:5173/\n")
    p_urls.add_run("• Interfaz de Servicios API (Backend): ").bold = True
    p_urls.add_run("http://localhost:4000/api\n")
    p_urls.add_run("• Documentación Interactiva Swagger: ").bold = True
    p_urls.add_run("http://localhost:4000/api/docs\n")
    p_urls.add_run("• Motor de Base de Datos: ").bold = True
    p_urls.add_run("PostgreSQL 16 en puerto 5432 (Contenedor mercapleno-postgres)")

    add_section_header("2.3. Matriz de Control de Acceso por Roles (RBAC)", level=2)
    doc.add_paragraph(
        "El sistema implementa un modelo estricto de Control de Acceso Basado en Roles (Role-Based Access Control). "
        "La siguiente matriz define las facultades operativas de cada perfil:"
    )
    
    t_rbac = doc.add_table(rows=8, cols=4)
    t_rbac_data = [
        ["Módulo / Funcionalidad", "Rol 1: Administrador", "Rol 2: Empleado", "Rol 3: Cliente"],
        ["Catálogo y Exploración de Productos", "Permitido", "Permitido", "Permitido"],
        ["Carrito de Compras y Pedidos", "Permitido", "Permitido", "Permitido"],
        ["Dashboard de Operaciones (/usuarioC)", "Permitido", "Permitido", "Denegado (403)"],
        ["Registro de Movimientos de Inventario", "Permitido (Total)", "Permitido (Total)", "Denegado (403)"],
        ["CRUD Catálogo Comercial (/products/admin)", "Permitido (Crear/Editar)", "Denegado (Solo lectura)", "Denegado (403)"],
        ["Administración de Proveedores", "Permitido", "Denegado (403)", "Denegado (403)"],
        ["Gestión de Usuarios y Roles (/admin/users)", "Permitido", "Denegado (403)", "Denegado (403)"],
    ]
    for r_idx, row in enumerate(t_rbac_data):
        for c_idx, val in enumerate(row):
            t_rbac.cell(r_idx, c_idx).paragraphs[0].add_run(val)
    style_table(t_rbac)

    # -------------------------------------------------------------
    # CAPÍTULO 3: SEGURIDAD, REGISTRO Y AUTENTICACIÓN
    # -------------------------------------------------------------
    add_section_header("3. Módulo 1: Seguridad, Registro y Autenticación (Épica EP-001)", level=1)
    
    doc.add_paragraph(
        "La seguridad en Mercapleno garantiza que todas las transacciones, datos personales e historiales comerciales "
        "se encuentren protegidos bajo estándares de cifrado robusto (Bcrypt), tokens de sesión (JWT) y verificación de doble factor."
    )

    add_section_header("3.1. Procedimiento de Registro de Nuevos Clientes", level=2)
    doc.add_paragraph(
        "Para darse de alta en la plataforma como nuevo cliente, siga estos pasos:"
    )
    p_reg = doc.add_paragraph()
    p_reg.add_run("Paso 1: ").bold = True
    p_reg.add_run("Acceda a la URL principal ")
    p_reg.add_run("http://localhost:5173/").bold = True
    p_reg.add_run(". Al cargar la vista de inicio de sesión, haga clic en el enlace inferior «¿No tiene cuenta? Regístrese».\n")
    p_reg.add_run("Paso 2: ").bold = True
    p_reg.add_run("Diligencie el formulario con sus datos obligatorios:\n")
    p_reg.add_run("  • Nombres y Apellidos completos.\n")
    p_reg.add_run("  • Tipo de Identificación (Cédula de ciudadanía, Tarjeta de identidad, Cédula de extranjería, Pasaporte o NIT).\n")
    p_reg.add_run("  • Número de Documento (único en el sistema; si ya existe, se notificará de inmediato).\n")
    p_reg.add_run("  • Correo Electrónico (será su identificador principal y canal de verificación).\n")
    p_reg.add_run("  • Fecha de Nacimiento y Dirección de Residencia.\n")
    p_reg.add_run("  • Contraseña Segura (mínimo 8 caracteres, combinación de mayúsculas, minúsculas, números y caracteres especiales).\n")
    p_reg.add_run("Paso 3: ").bold = True
    p_reg.add_run("Haga clic en el botón «Registrarse». El sistema procesará la solicitud y enviará un código numérico a su correo electrónico.")

    add_callout(doc,
                "El código de verificación tiene una vigencia máxima de 15 minutos (TTL). "
                "Si el código expira, utilice la opción «Reenviar código de verificación» en la pantalla de validación.",
                title="REGLA DE SEGURIDAD - CÓDIGOS TEMPORALES", callout_type="SECURITY")

    add_section_header("3.2. Activación de Cuenta mediante Código de Correo (OTP)", level=2)
    doc.add_paragraph(
        "Tras completar el formulario, será redirigido a la pantalla /verificar:\n"
        "1. Revise la bandeja de entrada o carpeta de spam del correo registrado.\n"
        "2. Copie el código alfanumérico o numérico de 6 dígitos recibido.\n"
        "3. Ingréselo en el campo de texto y presione «Validar Cuenta».\n"
        "4. Una vez validado, el sistema marcará su cuenta como verificada (email_verified: true) y podrá iniciar sesión inmediatamente."
    )

    add_section_header("3.3. Inicio de Sesión Estándar (Clientes)", level=2)
    doc.add_paragraph(
        "1. Ingrese a /login.\n"
        "2. Ingrese su correo electrónico y contraseña.\n"
        "3. Presione «Iniciar Sesión». Al tratarse de un cliente (Rol 3), el sistema lo redirigirá directamente al Catálogo Comercial (/catalogo)."
    )

    add_section_header("3.4. Inicio de Sesión con Doble Factor de Autenticación - 2FA (Administradores y Empleados)", level=2)
    doc.add_paragraph(
        "Para proteger la información contable y administrativa, las cuentas de Administrador (Rol 1) y Empleado (Rol 2) "
        "cuentan con un protocolo obligatorio de Doble Factor (2FA):\n"
        "1. Ingrese su correo y contraseña en /login.\n"
        "2. Al detectar un perfil administrativo, el sistema generará un Token Temporal de Acceso (pendingToken) y remitirá un código 2FA exclusivo a su correo corporativo.\n"
        "3. El sistema mostrará la interfaz modal «Verificación de Seguridad 2FA».\n"
        "4. Digite el código recibido (vigencia de 10 minutos) y presione «Verificar y Entrar».\n"
        "5. El sistema emitirá la cookie segura con el Access Token JWT definitivo y lo dirigirá al Dashboard de Operaciones (/usuarioC)."
    )

    add_section_header("3.5. Recuperación de Contraseña Olvidada", level=2)
    doc.add_paragraph(
        "En caso de extravío de sus credenciales:\n"
        "1. En la pantalla de Login, presione «¿Olvidó su contraseña?».\n"
        "2. Digite su correo electrónico registrado y presione «Enviar Código de Recuperación».\n"
        "3. Recibirá un mensaje con un código de restablecimiento temporal (válido por 15 minutos).\n"
        "4. En la pantalla /recuperar, ingrese el código recibido, digite su nueva contraseña y confírmela.\n"
        "5. Al hacer clic en «Restablecer Contraseña», el sistema actualizará el hash criptográfico y podrá acceder con su nueva clave."
    )

    # -------------------------------------------------------------
    # CAPÍTULO 4: CATÁLOGO Y TIENDA VIRTUAL
    # -------------------------------------------------------------
    add_section_header("4. Módulo 2: Catálogo de Productos y Tienda Virtual (Épica EP-003)", level=1)
    
    doc.add_paragraph(
        "El módulo de catálogo comercial (/catalogo) constituye la vitrina digital de Mercapleno, permitiendo "
        "a los clientes explorar la oferta de productos con información fidedigna de precios, categorías y disponibilidad en inventario."
    )

    add_section_header("4.1. Navegación y Componentes de la Ficha de Producto", level=2)
    doc.add_paragraph(
        "Cada tarjeta de producto en el catálogo presenta la siguiente información estructurada:\n"
        "• Imagen descriptiva de alta resolución.\n"
        "• Nombre comercial y descripción detallada del artículo.\n"
        "• Etiqueta de Categoría (Abarrotes, Lácteos, Cárnicos, Bebidas, Panadería, Frutas/Verduras, Aseo, etc.).\n"
        "• Precio unitario de venta en pesos colombianos ($ COP).\n"
        "• Indicador de Existencias (Stock actual en tiempo real).\n"
        "• Botón de Acción «Agregar al Carrito» con selector de cantidad."
    )

    add_section_header("4.2. Búsqueda y Filtros Avanzados", level=2)
    doc.add_paragraph(
        "La interfaz dispone de herramientas dinámicas de filtrado:\n"
        "1. Barra de Búsqueda por Texto: Filtre instantáneamente escribiendo el nombre del producto (ej. «Arroz», «Leche», «Ariel»).\n"
        "2. Selector de Categorías: Filtre los artículos haciendo clic en la categoría deseada en el panel superior.\n"
        "3. Filtro por Rango de Precios: Ordene los productos de menor a mayor precio o defina topes mínimos y máximos."
    )

    # -------------------------------------------------------------
    # CAPÍTULO 5: CARRITO DE COMPRAS Y PEDIDOS
    # -------------------------------------------------------------
    add_section_header("5. Módulo 3: Carrito de Compras y Proceso de Venta (Épica EP-003)", level=1)
    
    doc.add_paragraph(
        "El carrito de compras (/cart) permite consolidar la selección de artículos, verificar la disponibilidad física "
        "en bodega, liquidar impuestos y registrar formalmente la orden de compra."
    )

    add_section_header("5.1. Gestión de Ítems en el Carrito", level=2)
    p_cart_ops = doc.add_paragraph()
    p_cart_ops.add_run("• Adición de Productos: ").bold = True
    p_cart_ops.add_run("Al pulsar «Agregar al Carrito», el sistema registra el ítem en la base de datos PostgreSQL vinculándolo al carrito activo del usuario.\n")
    p_cart_ops.add_run("• Modificación de Cantidad: ").bold = True
    p_cart_ops.add_run("Utilice los botones [+] y [-] para ajustar las unidades. El sistema recalcula el subtotal en tiempo real.\n")
    p_cart_ops.add_run("• Validación de Stock: ").bold = True
    p_cart_ops.add_run("El sistema impide seleccionar cantidades superiores a las existencias registradas en la tabla stock_actual. Si un producto cuenta con stock bajo, se desplegará una alerta visual preventiva.\n")
    p_cart_ops.add_run("• Eliminación de Artículos: ").bold = True
    p_cart_ops.add_run("Haga clic en el ícono de papelera para remover un ítem individual o en «Vaciar Carrito» para limpiar la selección.")

    add_section_header("5.2. Liquidación Económica del Pedido", level=2)
    doc.add_paragraph(
        "El resumen de compra presenta el desglose financiero exacto de la transacción:"
    )
    t_liq = doc.add_table(rows=4, cols=2)
    t_liq_data = [
        ["Concepto Liquidado", "Fórmula / Regla de Cálculo"],
        ["Subtotal", "Suma de los productos seleccionados (Cantidad × Precio Snapshot)."],
        ["Impuesto al Valor Agregado (IVA)", "Cálculo legal del 19% sobre el subtotal gravable."],
        ["Total a Pagar", "Subtotal + IVA liquidado."]
    ]
    for r_idx, row in enumerate(t_liq_data):
        for c_idx, val in enumerate(row):
            t_liq.cell(r_idx, c_idx).paragraphs[0].add_run(val)
    style_table(t_liq)

    add_section_header("5.3. Métodos de Pago Disponibles", level=2)
    doc.add_paragraph(
        "Mercapleno soporta 6 métodos de pago integrados:\n"
        "• M1: Efectivo (pago contra entrega en caja o domicilio).\n"
        "• M2: Tarjeta de Crédito (Visa, Mastercard, American Express).\n"
        "• M3: Tarjeta de Débito (Cuentas de ahorro/corriente).\n"
        "• M4: Transferencia Bancaria (PSE / Botón bancario).\n"
        "• M5: Nequi (Pago móvil mediante notificación o QR).\n"
        "• M6: Daviplata (Billetera digital Davivienda)."
    )

    add_section_header("5.4. Confirmación y Emisión de Ticket Digital (/ticket)", level=2)
    doc.add_paragraph(
        "Al pulsar «Confirmar y Finalizar Pedido»:\n"
        "1. El sistema ejecuta una transacción ACID en PostgreSQL (descuenta el stock físico, registra la salida en salida_productos, crea la orden en venta y los ítems en venta_productos).\n"
        "2. El estado del carrito cambia de active a ordered.\n"
        "3. El usuario es redirigido a la vista de Ticket Digital (/ticket), donde puede visualizar el número de comprobante, fecha, productos adquiridos, método de pago y opción de impresión."
    )

    # -------------------------------------------------------------
    # CAPÍTULO 6: DASHBOARD DE OPERACIONES
    # -------------------------------------------------------------
    add_section_header("6. Módulo 4: Dashboard y Panel de Operaciones (Roles 1 y 2)", level=1)
    
    doc.add_paragraph(
        "El Dashboard (/usuarioC) es el centro neurálgico para la administración de la tienda. Desde este panel, "
        "los usuarios con rol de Administrador y Empleado pueden monitorear el estado general del negocio y acceder a los submódulos operativos."
    )
    
    doc.add_paragraph(
        "Accesos rápidos disponibles en el Dashboard:\n"
        "• Tarjetas de Resumen: Métricas de ventas del día, productos con stock crítico y órdenes recientes.\n"
        "• Botón «Inventario y Movimientos»: Acceso a la gestión de entradas, salidas y existencias.\n"
        "• Botón «Gestión de Productos» (Solo Admin): Acceso al CRUD comercial del catálogo.\n"
        "• Botón «Proveedores» (Solo Admin): Directorio y altas de proveedores.\n"
        "• Botón «Gestión de Usuarios» (Solo Admin): Control de cuentas y perfiles.\n"
        "• Botón «Estadísticas y Reportes»: Módulo analítico y exportación PDF."
    )

    # -------------------------------------------------------------
    # CAPÍTULO 7: GESTIÓN DE INVENTARIO Y MOVIMIENTOS
    # -------------------------------------------------------------
    add_section_header("7. Módulo 5: Gestión de Inventario y Registro de Movimientos (Épica EP-004)", level=1)
    
    doc.add_paragraph(
        "La correcta administración del inventario garantiza la continuidad de la operación comercial y previene "
        "quiebres de stock o mermas por vencimiento."
    )

    add_section_header("7.1. Registro de Entrada de Mercancía (Compras a Proveedores)", level=2)
    doc.add_paragraph(
        "Ruta: /products/employee -> Pestaña «Registrar Entrada»\n"
        "1. Seleccione el Producto receptor del lote.\n"
        "2. Ingrese la Cantidad de unidades que ingresan a bodega.\n"
        "3. Seleccione el Tipo de Documento Soporte (FA: Factura, CC: Cuenta de cobro, PE: Pedido).\n"
        "4. Indique el Costo Unitario de compra y la Fecha de Vencimiento del lote.\n"
        "5. Ingrese las Observaciones pertinentes (ej. «Llegada de camión proveedor #3»).\n"
        "6. Presione «Guardar Entrada». El sistema sumará automáticamente las existencias al stock actual y registrará el movimiento en entrada_productos."
    )

    add_section_header("7.2. Registro de Salidas Operativas y Mermas", level=2)
    doc.add_paragraph(
        "Ruta: /products/employee -> Pestaña «Registrar Salida»\n"
        "1. Seleccione el Producto a retirar.\n"
        "2. Ingrese la Cantidad a descargar.\n"
        "3. Seleccione el Motivo del movimiento (Merma, Producto averiado, Traslado interno o Ajuste de inventario).\n"
        "4. Presione «Registrar Salida». Las unidades serán descontadas de stock_actual y documentadas en salida_productos."
    )

    add_section_header("7.3. Semáforo de Alertas de Stock Bajo", level=2)
    doc.add_paragraph(
        "El sistema calcula dinámicamente el estado del inventario aplicando un umbral paramétrico (LOW_STOCK_THRESHOLD = 5 unidades):\n"
        "• Verde (Óptimo): Stock > 10 unidades.\n"
        "• Amarillo (Alerta preventiva): Stock entre 6 y 10 unidades.\n"
        "• Rojo (Crítico / Bajo Stock): Stock <= 5 unidades. Se sugiere realizar pedido inmediato al proveedor."
    )

    # -------------------------------------------------------------
    # CAPÍTULO 8: ADMINISTRACIÓN DEL CATÁLOGO DE PRODUCTOS
    # -------------------------------------------------------------
    add_section_header("8. Módulo 6: Administración del Catálogo de Productos (Épica EP-004 - Admin)", level=1)
    
    doc.add_paragraph(
        "Ruta: /products/admin (Exclusivo para Rol 1: Administrador).\n"
        "Permite gestionar el ciclo de vida completo de los artículos comerciales."
    )

    add_section_header("8.1. Alta de un Nuevo Producto", level=2)
    doc.add_paragraph(
        "1. Presione el botón «+ Nuevo Producto».\n"
        "2. Diligencie el Nombre del producto (ej. «Arroz Diana 1000g»).\n"
        "3. Asigne la Categoría correspondiente del selector.\n"
        "4. Seleccione el Proveedor responsable del suministro.\n"
        "5. Fije el Precio de Venta al público en $ COP.\n"
        "6. Ingrese la Descripción comercial del artículo.\n"
        "7. Ingrese la URL de la Imagen oficial o cargue el archivo multimedia.\n"
        "8. Seleccione el Estado inicial (Disponible).\n"
        "9. Presione «Guardar Producto». El artículo quedará disponible de inmediato en el catálogo público."
    )

    add_section_header("8.2. Modificación y Deshabilitación (Soft Delete)", level=2)
    doc.add_paragraph(
        "• Edición: En la tabla de productos, presione el ícono de lápiz en la fila del artículo, actualice los campos y guarde los cambios.\n"
        "• Deshabilitación: Si un producto deja de comercializarse temporalmente, cambie su estado a Deshabilitado. "
        "El producto se ocultará de la tienda de clientes pero preservará toda la integridad referencial de los reportes y ventas históricas."
    )

    # -------------------------------------------------------------
    # CAPÍTULO 9: ADMINISTRACIÓN DE PROVEEDORES
    # -------------------------------------------------------------
    add_section_header("9. Módulo 7: Gestión de Proveedores (Rol Administrador)", level=1)
    
    doc.add_paragraph(
        "Ruta: /admin/proveedores\n"
        "El módulo de proveedores permite mantener la relación comercial y los canales de abastecimiento actualizados."
    )
    doc.add_paragraph(
        "Funcionalidades disponibles:\n"
        "1. Listado y Búsqueda: Consulte proveedores por nombre, apellido o número de teléfono.\n"
        "2. Crear Proveedor: Registre el nombre de la empresa/contacto y su teléfono móvil de 10 dígitos.\n"
        "3. Habilitar / Deshabilitar: Active o inactive proveedores según el estado de la relación contractual sin perder el historial de compras."
    )

    # -------------------------------------------------------------
    # CAPÍTULO 10: ADMINISTRACIÓN DE USUARIOS
    # -------------------------------------------------------------
    add_section_header("10. Módulo 8: Administración de Usuarios y Roles (Épica EP-002 - Admin)", level=1)
    
    doc.add_paragraph(
        "Ruta: /admin/users (Exclusivo para Rol 1: Administrador).\n"
        "Permite gobernar las identidades y niveles de acceso de todo el personal de Mercapleno."
    )
    doc.add_paragraph(
        "Procedimientos principales:\n"
        "• Consulta de Cuentas: Búsqueda rápida por nombre, correo, número de documento o rol asignado.\n"
        "• Creación de Cuentas Administrativas: Formulario para dar de alta nuevos Empleados o Administradores asignando credenciales iniciales.\n"
        "• Edición de Perfiles: Modificación de direcciones, teléfonos, correos y cambio de rol (ej. promover un Empleado a Administrador).\n"
        "• Eliminación Segura: Supresión de cuentas de usuario cuando ya no forman parte de la organización."
    )

    # -------------------------------------------------------------
    # CAPÍTULO 11: REPORTES, ANALÍTICA Y EXPORTACIÓN PDF
    # -------------------------------------------------------------
    add_section_header("11. Módulo 9: Reportes, Analítica y Exportación PDF (Épica EP-005)", level=1)
    
    doc.add_paragraph(
        "Ruta: /estadisticas (Acceso para Roles 1 y 2).\n"
        "El módulo analítico procesa los registros transaccionales para generar métricas ejecutivas de rendimiento comercial."
    )

    add_section_header("11.1. Indicadores Clave de Desempeño (KPIs)", level=2)
    doc.add_paragraph(
        "El panel presenta tres visualizaciones principales:\n"
        "1. Ventas por Mes: Historial cronológico de ingresos totales facturados mes a mes.\n"
        "2. Top Productos Más Vendidos: Gráfico y ranking de los artículos con mayor rotación y volumen de ventas.\n"
        "3. Resumen Ejecutivo: Total de transacciones procesadas, ticket promedio de venta y valor total recaudado."
    )

    add_section_header("11.2. Descarga de Reportes en Formato PDF", level=2)
    doc.add_paragraph(
        "1. En la parte superior de la vista de estadísticas, localice el botón «📄 Descargar Reporte PDF».\n"
        "2. El servidor backend procesará las tablas de venta en tiempo real y compilará un archivo PDF formal con membrete de Mercapleno, tablas de resumen y fecha de emisión.\n"
        "3. El archivo se descargará automáticamente en el navegador para su archivo, impresión o presentación gerencial."
    )

    # -------------------------------------------------------------
    # CAPÍTULO 12: SOLUCIÓN DE PROBLEMAS Y FAQ
    # -------------------------------------------------------------
    add_section_header("12. Guía de Solución de Problemas (Troubleshooting) y Preguntas Frecuentes", level=1)
    
    doc.add_paragraph(
        "A continuación se detallan los incidentes operativos más comunes y su procedimiento de resolución inmediata:"
    )

    t_faq = doc.add_table(rows=7, cols=3)
    t_faq_data = [
        ["Incidente / Mensaje de Error", "Posible Causa", "Solución Recomendada"],
        ["«Correo o contraseña incorrectos»", "Credenciales erradas o cuenta no registrada.", "Verifique mayúsculas/minúsculas. Si olvidó su clave, utilice la opción «Recuperar contraseña»."],
        ["«Código de verificación expirado»", "Transcurrieron más de 15 minutos desde el envío.", "Presione «Reenviar código» en la pantalla de verificación para obtener uno nuevo."],
        ["«No recibí el código 2FA»", "El correo se encuentra en carpeta de spam o SMTP saturado.", "Revise carpetas de Correo no deseado / Spam. Espere 1 minuto y solicite un reenvío."],
        ["«Stock insuficiente al agregar al carrito»", "La cantidad solicitada supera las existencias en bodega.", "Ajuste la cantidad en el selector o espere a que el administrador registre una entrada de mercancía."],
        ["«Acceso Denegado (403)»", "Su rol actual no tiene permisos para ingresar a esa vista.", "Inicie sesión con una cuenta que disponga del rol requerido (Administrador o Empleado)."],
        ["«Error 500 / Conexión rechazada»", "El contenedor de backend o PostgreSQL no está activo.", "Verifique que Docker esté en ejecución y los contenedores mercapleno-backend y mercapleno-postgres estén en estado Up."]
    ]
    for r_idx, row in enumerate(t_faq_data):
        for c_idx, val in enumerate(row):
            t_faq.cell(r_idx, c_idx).paragraphs[0].add_run(val)
    style_table(t_faq)

    # -------------------------------------------------------------
    # CAPÍTULO 13: GLOSARIO TÉCNICO Y DE NEGOCIO
    # -------------------------------------------------------------
    add_section_header("13. Glosario Técnico y de Términos del Negocio", level=1)
    
    glossary = [
        ("2FA (Two-Factor Authentication):", "Mecanismo de seguridad que exige dos formas de autenticación: la contraseña del usuario y un código temporal enviado por correo."),
        ("Bcrypt:", "Algoritmo de derivación de claves criptográficas y hashing utilizado para almacenar contraseñas de forma irreversible."),
        ("JWT (JSON Web Token):", "Estándar abierto para la transmisión segura de información de identidad firmada digitalmente entre el cliente y el servidor."),
        ("OTP (One-Time Password):", "Contraseña de un solo uso con vigencia limitada en el tiempo (TTL)."),
        ("Price Snapshot:", "Captura del precio unitario de un producto en el instante exacto en que se añade al carrito, protegiendo al cliente ante cambios de precio posteriores."),
        ("RBAC (Role-Based Access Control):", "Método de regulación de acceso a los recursos de un sistema basado en los roles individuales asignados a cada usuario."),
        ("Soft Delete:", "Técnica de desactivación de registros en base de datos mediante cambio de estado sin eliminación física, preservando la integridad histórica."),
        ("Stock Actual:", "Cantidad total de unidades físicas de un artículo disponibles en almacén para su comercialización inmediata."),
        ("Ticket Digital:", "Comprobante electrónico emitido tras finalizar una compra que detalla los productos adquiridos, montos, impuestos y método de pago.")
    ]
    
    for term, definition in glossary:
        p_g = doc.add_paragraph()
        p_g.paragraph_format.space_before = Pt(2)
        p_g.paragraph_format.space_after = Pt(4)
        r_term = p_g.add_run(term + " ")
        r_term.bold = True
        r_term.font.color.rgb = RGBColor(30, 58, 138)
        r_def = p_g.add_run(definition)
        r_def.font.color.rgb = RGBColor(71, 85, 105)

    # Save documents
    target_path1 = r"C:\Users\jefer\Downloads\Manual de Usuario - Mercapleno Profesional.docx"
    target_path2 = r"C:\Users\jefer\Downloads\Manual de usuario.docx"
    
    doc.save(target_path1)
    print(f"Documento guardado exitosamente en: {target_path1}")
    
    try:
        doc.save(target_path2)
        print(f"Borrador actualizado exitosamente en: {target_path2}")
    except Exception as e:
        print(f"Aviso al sobreescribir borrador: {e}")

if __name__ == "__main__":
    build_manual()

