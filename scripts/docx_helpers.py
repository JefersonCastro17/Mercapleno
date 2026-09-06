import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import os

def create_element(name):
    return OxmlElement(name)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='003366')
    """
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b_elm = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val", "single")}" w:sz="{edge_data.get("sz", "4")}" w:space="0" w:color="{edge_data.get("color", "auto")}"/>')
            tcBorders.append(b_elm)
    tcPr.append(tcBorders)

def add_callout(doc, text, title="NOTA IMPORTANTE", callout_type="NOTE"):
    colors = {
        "NOTE": {"bg": "F0F7FF", "border": "2563EB", "title_color": RGBColor(37, 99, 235)},
        "WARNING": {"bg": "FFFBEB", "border": "D97706", "title_color": RGBColor(217, 119, 6)},
        "TIP": {"bg": "F0FDF4", "border": "16A34A", "title_color": RGBColor(22, 163, 74)},
        "SECURITY": {"bg": "FEF2F2", "border": "DC2626", "title_color": RGBColor(220, 38, 38)},
    }
    cfg = colors.get(callout_type, colors["NOTE"])
    
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, cfg["bg"])
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    set_cell_border(cell, 
                    left={"val": "single", "sz": "24", "color": cfg["border"]},
                    top={"val": "none", "sz": "0", "color": "auto"},
                    bottom={"val": "none", "sz": "0", "color": "auto"},
                    right={"val": "none", "sz": "0", "color": "auto"})
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}")
    run_t.bold = True
    run_t.font.name = "Segoe UI"
    run_t.font.size = Pt(10)
    run_t.font.color.rgb = cfg["title_color"]
    
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    run_body = p2.add_run(text)
    run_body.font.name = "Segoe UI"
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_table(table, header_bg="1E3A8A", alt_bg="F8FAFC"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        
        if i == 0:
            # Header row
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            for cell in row.cells:
                set_cell_background(cell, header_bg)
                set_cell_margins(cell, top=140, bottom=140, left=150, right=150)
                set_cell_border(cell, 
                                bottom={"val": "single", "sz": "12", "color": "0F172A"},
                                top={"val": "single", "sz": "4", "color": "CBD5E1"},
                                left={"val": "single", "sz": "4", "color": "CBD5E1"},
                                right={"val": "single", "sz": "4", "color": "CBD5E1"})
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.bold = True
                        r.font.name = "Segoe UI"
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(255, 255, 255)
        else:
            bg = alt_bg if i % 2 == 1 else "FFFFFF"
            for cell in row.cells:
                set_cell_background(cell, bg)
                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                set_cell_border(cell, 
                                bottom={"val": "single", "sz": "4", "color": "E2E8F0"},
                                top={"val": "single", "sz": "4", "color": "E2E8F0"},
                                left={"val": "single", "sz": "4", "color": "E2E8F0"},
                                right={"val": "single", "sz": "4", "color": "E2E8F0"})
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.name = "Segoe UI"
                        r.font.size = Pt(9)
                        r.font.color.rgb = RGBColor(51, 65, 85)

print("Helper functions ready")

